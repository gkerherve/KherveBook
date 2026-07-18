"""Readable previews for structured attachment formats.

A File cell shows a preview of each attached file. Plain text and images
are handled in ``filecell`` directly; this module reads the binary
scientific formats in the Kherve family so their *contents* — every
sheet, every core level — can be browsed, not just labelled "binary":

* ``.xlsx`` / ``.xlsm`` — Excel workbook (openpyxl): one part per sheet,
  each a small cell grid.
* ``.ksheet``           — KherveSheet workbook (HDF5): one part per sheet.
* ``.kfit``             — KherveFitting project (HDF5 + a zlib JSON blob):
  one part per core level (B.E. range, point count, fitted peaks). The
  full core-level list is read from the project JSON, which is
  authoritative — the HDF5 ``core_levels`` group may hold only the level
  currently displayed.

``describe()`` returns ``{"header": str, "parts": [{"name", "text"}, …]}``
so the cell can offer a selector to read each part. Every reader is
best-effort and optional-dependency-safe: a missing library or an
unexpected layout returns ``None`` and the cell falls back to the
generic "binary file — kept as-is" note.

Copyright (C) 2026 Gwilherm Kerherve

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.
"""

import io
import json
import zlib
from pathlib import Path

#: preview budget — keep each part's snippet compact in the cell.
_MAX_ROWS = 8
_MAX_COLS = 8
_CELL_W = 13
_MAX_PEAKS = 8


def _text(v) -> str:
    if isinstance(v, bytes):
        return v.decode("utf-8", "replace")
    return "" if v is None else str(v)


def _cell(v) -> str:
    return "" if v is None else str(v)


def _grid_text(rows) -> str:
    """A list of row-lists -> an aligned, column-padded snippet."""
    out = []
    for row in rows[:_MAX_ROWS]:
        cells = [_cell(v)[:_CELL_W].ljust(_CELL_W) for v in row[:_MAX_COLS]]
        out.append("  ".join(cells).rstrip())
    return "\n".join(out).rstrip() or "(empty)"


def describe(name: str, data: bytes):
    """Return {"header", "parts":[{"name","text"}]} for a structured file,
    or None if this module can't read it."""
    ext = Path(name).suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(data)
        if ext in (".ksheet", ".kfit"):
            return _hdf5(data)
    except Exception:
        return None
    return None


# -- Excel -----------------------------------------------------------------

def _xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        parts = []
        for nm in names:
            ws = wb[nm]
            rows = []
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= _MAX_ROWS:
                    break
                rows.append(list(row))
            dims = getattr(ws, "calculate_dimension", lambda: "")()
            parts.append({"name": nm,
                          "text": f"[{nm}] {dims}\n" + _grid_text(rows)})
        n = len(names)
        header = f"Excel workbook · {n} sheet{'s' if n != 1 else ''}"
        return {"header": header, "parts": parts}
    finally:
        wb.close()


# -- HDF5 (KherveSheet .ksheet and KherveFitting .kfit) --------------------

def _hdf5(data: bytes):
    import h5py
    with h5py.File(io.BytesIO(data), "r") as f:
        fmt = _text(f.attrs.get("format", "")).lower()
        if fmt == "ksheet":
            return _ksheet(f)
        if fmt == "kfitting":
            return _kfit(f)
    return None


def _ksheet(f):
    count = int(f.attrs.get("sheet_count", 0))
    parts = []
    for si in range(count):
        g = f.get(f"sheet_{si}")
        if g is None:
            continue
        nm = _text(g.attrs.get("name", f"Sheet {si + 1}"))
        rows = int(g.attrs.get("rows", 0))
        cols = int(g.attrs.get("cols", 0))
        charts = sum(1 for k in g if str(k).startswith("chart_"))
        grid = []
        if "cells" in g and cols:
            flat = g["cells"][:]
            for r in range(min(rows, _MAX_ROWS)):
                grid.append([_text(flat[r * cols + c])
                             for c in range(min(cols, _MAX_COLS))
                             if r * cols + c < len(flat)])
        tag = f"{nm}  ({rows}×{cols}"
        tag += f", {charts} chart{'s' if charts != 1 else ''})" if charts \
            else ")"
        parts.append({"name": nm, "text": tag + "\n" + _grid_text(grid)})
    n = len(parts)
    return {"header": f"KherveSheet workbook · {n} sheet{'s' if n != 1 else ''}",
            "parts": parts}


def _kfit(f):
    js = _kfit_json(f)
    levels = js.get("Core levels") if isinstance(js, dict) else None
    names = []
    if isinstance(levels, dict) and levels:
        names = list(levels.keys())          # authoritative full list
    else:                                    # fall back to the HDF5 group
        cl = f.get("core_levels")
        if cl is not None:
            def _order(s):
                parts = str(s).split("_")
                return (0, int(parts[1])) if parts[-1].isdigit() \
                    else (1, str(s))
            for key in sorted(cl.keys(), key=_order):
                names.append(_text(cl[key].attrs.get("name", key)))
            levels = None

    parts = [{"name": nm,
              "text": _kfit_level_text(nm, levels.get(nm) if levels else None)}
             for nm in names]
    sample = _kfit_sample_name(js)
    n = len(names)
    header = f"KherveFitting project · {n} core level{'s' if n != 1 else ''}"
    if sample:
        header += f"  ·  {sample}"
    return {"header": header, "parts": parts}


def _kfit_level_text(name: str, lvl) -> str:
    if not isinstance(lvl, dict):
        return name
    be = lvl.get("B.E.") or []
    lines = []
    if be:
        lines.append(f"{name} — {len(be)} points, "
                     f"B.E. {min(be):.1f}–{max(be):.1f} eV")
    else:
        lines.append(name)
    peaks = (lvl.get("Fitting") or {}).get("Peaks")
    if isinstance(peaks, dict) and peaks:
        lines.append(f"{len(peaks)} fitted peak"
                     f"{'s' if len(peaks) != 1 else ''}:")
        for pname, pk in list(peaks.items())[:_MAX_PEAKS]:
            if isinstance(pk, dict):
                pos = pk.get("Position")
                fwhm = pk.get("FWHM")
                area = pk.get("Area")
                bits = [f"pos {pos:.2f}" if isinstance(pos, (int, float))
                        else "", f"FWHM {fwhm:.2f}"
                        if isinstance(fwhm, (int, float)) else "",
                        f"area {area:.0f}" if isinstance(area, (int, float))
                        else ""]
                lines.append(f"  {pname:<14} "
                             + "  ".join(b for b in bits if b))
        if len(peaks) > _MAX_PEAKS:
            lines.append(f"  … (+{len(peaks) - _MAX_PEAKS} more)")
    else:
        lines.append("(no fitted peaks)")
    return "\n".join(lines)


def _kfit_json(f):
    ds = f.get("project_json_gz")
    if ds is None:
        return {}
    try:
        return json.loads(zlib.decompress(bytes(ds[()])).decode("utf-8"))
    except Exception:
        return {}


def _kfit_sample_name(js):
    for key in ("SampleNames", "FilePath"):
        val = js.get(key) if isinstance(js, dict) else None
        if isinstance(val, list) and val:
            return Path(_text(val[0])).name
        if isinstance(val, str) and val:
            return Path(val).name
    return ""
