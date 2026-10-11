"""Thermo Avantage Excel exports, and the workbook round trip.

Split from `thermo` (file-size policy). Ported from KherveFittingPro's
``libraries/FileMenu/Avantage_Import.py`` (single file, .xls, and the
"multiple files (folder)" import), ``MRS_Import.
extract_acquisition_parameters``, the Peak Table -> Fitting import of
``add_peak_fitting_data_after_load`` and the Smart background it draws
(``Peak_Functions.BackgroundCalculations``), wx-free.

KherveFitting writes these (and AVG files) to an intermediate
``_Kfitting.xlsx`` and reads it back with ``ConfigFile.
add_core_level_Data``; `_sheet_core_level` reproduces that round trip
(two-decimal rounding, the Experimental Description column) and
`_xl_title` openpyxl's sheet-title de-duplication, without a file.

Copyright (C) 2026 Gwilherm Kerherve

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.
"""

import math
import os
import re

import numpy as np

# ---------------------------------------------------------------------------
# Shared: the workbook round trip KherveFitting does for AVG / Avantage
# ---------------------------------------------------------------------------

def _xl_title(names, value):
    """openpyxl's ``avoid_duplicate_name``: a sheet title already in the
    workbook (case-insensitive) gets the next integer appended."""
    if any(n.lower() == value.lower() for n in names):
        joined = ",".join(names)
        rx = re.compile(f'(?P<title>{re.escape(value)})(?P<count>\\d*),?', re.I)
        matches = rx.findall(joined)
        if matches:
            counts = [int(idx) for (_t, idx) in matches if idx.isdigit()]
            highest = max(counts) if counts else 0
            value = f"{value}{highest + 1}"
    return value


def _isna(v):
    if v is None:
        return True
    try:
        return isinstance(v, float) and math.isnan(v)
    except TypeError:
        return False


def _blank(v):
    """An Excel cell written with *v* reads back empty."""
    return v is None or (isinstance(v, str) and v == "")


def _sheet_core_level(name, rows, exp_rows):
    """``ConfigFile.add_core_level_Data`` on a sheet KherveFitting wrote.

    *rows* are the data rows under the header (columns A, B[, C, D]) and
    *exp_rows* the (parameter, value) pairs of the Experimental
    Description column. Column B lands in 'Raw Data', C in 'Corrected
    Data' (defaulting to B), D in 'Transmission' (defaulting to 1.0);
    everything is rounded to two decimals as the reader does.
    """
    be_values, raw_data, corrected, transmission = [], [], [], []
    for row in rows:
        row = list(row) + [None] * (4 - len(row))
        be_val, raw_val = row[0], row[1]
        if _blank(be_val) or _blank(raw_val) or _isna(be_val) \
                or _isna(raw_val):
            continue
        try:
            be_values.append(float(be_val))
            raw_data.append(float(raw_val))
            c = row[2]
            corrected.append(float(raw_val) if _blank(c) or _isna(c)
                             else float(c))
            t = row[3]
            transmission.append(1.0 if _blank(t) or _isna(t) else float(t))
        except (ValueError, TypeError):
            continue

    info = {}
    for key, value in exp_rows:
        if _blank(key) or not str(key).strip():
            continue
        info[str(key).strip()] = ("" if _blank(value)
                                  else str(value).strip())

    cl = {
        'B.E.': [float(f"{v:.2f}") for v in be_values],
        'Raw Data': [float(f"{v:.2f}") for v in raw_data],
        'Corrected Data': [float(f"{v:.2f}") for v in corrected],
        'Transmission': [float(f"{v:.2f}") for v in transmission],
        'Name': name,
    }
    if info:
        cl['ExperimentalInfo'] = info
    if be_values:
        cl['Background'] = {
            'Bkg Y': cl['Raw Data'],
            'Bkg Type': '',
            'Bkg Low': float(f"{min(be_values):.2f}"),
            'Bkg High': float(f"{max(be_values):.2f}"),
            'Bkg Offset Low': 0,
            'Bkg Offset High': 0,
        }
    return cl


# ---------------------------------------------------------------------------
# Avantage — Excel export (.xlsx / .xls) with a 'Titles' sheet
# ---------------------------------------------------------------------------

_VB_TERMS = ['valence', 'valence band', 'valence scan', 'vb scan', 'vb',
             'valence band scan']
_SPECTRUM_TERMS_XLSX = _VB_TERMS + ['val', 'fermi', 'fermi scan',
                                    'cut-off', 'cutoff', 'cut off']
_SPECTRUM_TERMS_XLS = _VB_TERMS + ['fermi', 'fermi scan',
                                   'cut-off', 'cutoff', 'cut off']


def _is_spectrum_sheet(name, terms):
    return ("Survey" in name or "Scan" in name
            or any(t in name.lower() for t in terms))


class _TempSheet:
    """The throwaway openpyxl sheet KherveFitting copies an .xls sheet's
    first 25 x 20 cells into, for ``extract_acquisition_parameters``."""

    class _Cell:
        def __init__(self, value):
            self.value = value

    def __init__(self, xls_sheet):
        self._cells = {}
        for r in range(min(25, xls_sheet.nrows)):
            for c in range(min(20, xls_sheet.ncols)):
                v = xls_sheet.cell_value(r, c)
                if v is not None and str(v).strip():
                    self._cells[(r + 1, c + 1)] = v
        self.max_row = max([r for r, _ in self._cells] or [1])
        self.max_column = max([c for _, c in self._cells] or [1])

    def cell(self, row, column):
        return self._Cell(self._cells.get((row, column)))


def extract_acquisition_parameters(sheet):
    """MRS_Import's reader of the 'Acquisition Parameters' block."""
    parameters = {}
    acq_col = None
    for label in ("Acquisition Parameters", "Parameters"):
        for row in range(1, min(5, sheet.max_row + 1)):
            for col in range(1, sheet.max_column + 1):
                v = sheet.cell(row=row, column=col).value
                if v and label in str(v):
                    acq_col = col
                    break
            if acq_col:
                break
        if acq_col:
            break
    if acq_col is None:
        return parameters

    for row in range(3, min(15, sheet.max_row + 1)):
        name = sheet.cell(row=row, column=acq_col).value
        value = sheet.cell(row=row, column=acq_col + 1).value
        value_j = sheet.cell(row=row, column=acq_col + 2).value
        if name is None:
            continue
        name = str(name).strip()
        if not name or name in ['Parameter', 'Parameters', '']:
            continue
        parts = []
        if value is not None and str(value).strip():
            parts.append(str(value).strip())
        if value_j is not None and str(value_j).strip():
            parts.append(str(value_j).strip())
        if parts:
            parameters[name] = " ".join(parts)
    return parameters


def extract_title_sheet_info(wb, sample_suffix=None):
    """'Titles' (or 'Titles A', ...) A1 / A3 / A5."""
    info = {}
    title_name = "Titles" if sample_suffix is None \
        else f"Titles {sample_suffix}"
    if title_name in wb.sheetnames:
        ws = wb[title_name]
        for row, key in ((1, 'Sample Info'), (3, 'Additional Info'),
                         (5, 'Notes')):
            v = ws.cell(row=row, column=1).value
            if v:
                info[key] = str(v).strip()
    return info


def extract_title_sheet_info_xls(wb_xls, sample_suffix=None):
    info = {}
    title_name = "Titles" if sample_suffix is None \
        else f"Titles {sample_suffix}"
    if title_name in wb_xls.sheet_names():
        ws = wb_xls.sheet_by_name(title_name)
        for row, key in ((0, 'Sample Info'), (2, 'Additional Info'),
                         (4, 'Notes')):
            try:
                v = ws.cell_value(row, 0)
                if v:
                    info[key] = str(v).strip()
            except Exception:
                pass
    return info


def convert_avantage_peak_name(avantage_name):
    """'Sr3d5 Sr-F2' -> 'Sr3d5/2 Sr-F2', 'Ti2p3 ?' -> 'Ti2p3/2 ?'."""
    full = {'p': {'3': '3/2', '1': '1/2'}, 'd': {'5': '5/2', '3': '3/2'},
            'f': {'7': '7/2', '5': '5/2'}}

    def repl(m):
        j = full.get(m.group(3), {}).get(m.group(4), m.group(4))
        return f"{m.group(1)}{m.group(2)}{m.group(3)}{j}{m.group(5)}"

    return re.sub(r'([A-Z][a-z]?)(\d+)([pdf])(\d)(\s|$)', repl,
                  avantage_name)


def extract_peak_fitting_from_peak_table(wb):
    """'Peak Fit Table : O1s Scan' blocks of the 'Peak Table' sheet ->
    {scan: [peak dict]}; Avantage ref letters L, M, N become A, B, C."""
    data = {}
    if "Peak Table" not in wb.sheetnames:
        return data
    sheet = wb["Peak Table"]

    def convert_refs(text, mapping):
        if not text or not isinstance(text, str):
            return text
        for av, kf in mapping.items():
            text = re.sub(rf'\b{av}\b', kf, text)
        return text

    def clean_range(text):
        if not text or not isinstance(text, str):
            return text
        if ':' in text and text.endswith(':00'):
            return text[:-3]
        return text

    current = None
    header_row = None
    letters = {}
    row_idx = 1
    while row_idx <= sheet.max_row:
        cell = sheet.cell(row=row_idx, column=1).value
        if cell and isinstance(cell, str) and "Peak Fit Table :" in cell:
            parts = cell.split(":")
            if len(parts) >= 2:
                current = parts[1].strip().split()[0].strip()
                data[current] = []
                header_row = row_idx + 1
                letters = {}
                row_idx += 2
                continue

        if current and header_row and row_idx > header_row:
            ref = sheet.cell(row=row_idx, column=1).value
            if not ref:
                row_idx += 1
                continue
            if "Peak Fit Table" in str(ref):
                current = None
                header_row = None
                continue

            name_value = sheet.cell(row=row_idx, column=2).value
            peak_be = sheet.cell(row=row_idx, column=3).value
            height = sheet.cell(row=row_idx, column=4).value
            area = sheet.cell(row=row_idx, column=6).value
            fwhm = sheet.cell(row=row_idx, column=8).value
            lg_mix = sheet.cell(row=row_idx, column=9).value

            ref_str = str(ref).strip()
            if ref_str not in letters:
                letters[ref_str] = chr(65 + len(letters))
            letter = letters[ref_str]

            constraints = {}
            if row_idx + 1 <= sheet.max_row:
                for key, col in (('position_constraint', 3),
                                 ('fwhm_constraint', 8),
                                 ('lg_constraint', 9)):
                    v = sheet.cell(row=row_idx + 1, column=col).value
                    if v:
                        constraints[key] = clean_range(
                            convert_refs(str(v).strip(), letters))
            try:
                raw_name = (str(name_value).strip() if name_value
                            else f"Peak_{letter}")
                data[current].append({
                    'ref': letter,
                    'name': convert_avantage_peak_name(raw_name),
                    'position': (float(peak_be) if peak_be
                                 and not isinstance(peak_be, str) else 0.0),
                    'height': float(height) if height else 0.0,
                    'area': float(area) if area else 0.0,
                    'fwhm': (float(fwhm) if fwhm
                             and not isinstance(fwhm, str) else 1.5),
                    'lg_ratio': (float(lg_mix) if lg_mix
                                 and not isinstance(lg_mix, str) else 20.0),
                    'position_constraint':
                        constraints.get('position_constraint'),
                    'fwhm_constraint': constraints.get('fwhm_constraint'),
                    'lg_constraint': constraints.get('lg_constraint'),
                })
            except (ValueError, TypeError):
                pass
            row_idx += 2
            continue
        row_idx += 1
    return data


# --- Smart background (Peak_Functions.BackgroundCalculations) -------------

def _round_sig(value, sig=6):
    try:
        v = float(value)
    except (TypeError, ValueError):
        return value
    if v == 0 or not math.isfinite(v):
        return v
    return round(v, sig - 1 - int(math.floor(math.log10(abs(v)))))


def _endpoint_average(x, y, point, num_points):
    if num_points <= 0:
        num_points = 1
    idx = int(np.argmin(np.abs(x - point)))
    half = num_points // 2
    ideal_start = idx - half
    ideal_end = idx + half + (1 if num_points % 2 == 1 else 0)
    n = len(y)
    if ideal_start < 0:
        start, end = 0, min(n, ideal_end - ideal_start)
    elif ideal_end > n:
        end, start = n, max(0, ideal_start - (ideal_end - n))
    else:
        start, end = ideal_start, ideal_end
    start, end = max(0, start), min(n, end)
    if end - start > num_points:
        excess = end - start - num_points
        start += excess // 2
        end -= excess - excess // 2
    sel = y[start:end]
    if len(sel) > 0:
        return float(_round_sig(np.mean(sel)))
    return float(_round_sig(y[idx]))


def _cumtrapz0(y, x):
    d = np.diff(x)
    return np.concatenate([[0.0], np.cumsum(d * (y[1:] + y[:-1]) / 2.0)])


def _shirley_calculate(x, y, tol=1e-5, maxit=10):
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    if len(x) < 2 or len(x) != len(y):
        return np.zeros_like(y)
    rev = y[0] < y[-1]
    if rev:
        x, y = x[::-1], y[::-1]
    i_hi, i_lo = y[0], y[-1]
    b = np.zeros_like(y)
    for _ in range(max(1, int(maxit))):
        signal = np.clip(y - i_lo - b, 0.0, None)
        q = _cumtrapz0(signal[::-1], x[::-1])[::-1]
        k = (i_hi - i_lo) / q[0] if q[0] != 0 else 0.0
        new = k * q
        change = float(np.mean((new - b) ** 2))
        b = new
        if tol > 0 and change < tol:
            break
    out = i_lo + b
    return out[::-1] if rev else out


def _shirley_background(x, y, start_offset, end_offset, max_iter=100,
                        tol=1e-2, num_points=5):
    x, y = np.asarray(x), np.asarray(y)
    dx = x[1] - x[0]
    x_pad = np.concatenate([[x[0] - dx], x, [x[-1] + dx]])
    y_start = _endpoint_average(x, y, x[0], num_points) + start_offset
    y_end = _endpoint_average(x, y, x[-1], num_points) + end_offset
    y_pad = np.concatenate([[y_start], y, [y_end]])
    scale2 = float(np.nanmean(np.square(y_pad))) if len(y_pad) else 0.0
    if np.isfinite(scale2) and scale2 > 0:
        tol = min(tol, 1e-8 * scale2)
    return _shirley_calculate(x_pad, y_pad, maxit=max_iter, tol=tol)[1:-1]


def _linear_background(x, y, start_offset, end_offset, num_points=5):
    y_start = _endpoint_average(x, y, x[0], num_points) + start_offset
    y_end = _endpoint_average(x, y, x[-1], num_points) + end_offset
    return np.linspace(y_start, y_end, len(y))


def _smart_background(x, y, offset_h, offset_l, num_points=5):
    shirley = _shirley_background(x, y, offset_h, offset_l,
                                  num_points=num_points)
    linear = _linear_background(x, y, offset_h, offset_l,
                                num_points=num_points)
    return shirley if y[0] > y[-1] else linear


def _adaptive_smart_background(x, y, x_range, previous, offset_h, offset_l,
                               num_points=5):
    previous = np.array(previous)
    mask = (x >= x_range[0]) & (x <= x_range[1])
    new = np.copy(previous)
    new[mask] = _smart_background(x[mask], y[mask], offset_h, offset_l,
                                  num_points)
    return new


def add_peak_fitting(levels, peak_fitting_by_scan):
    """``add_peak_fitting_data_after_load``: the Peak Table's peaks become
    each matching sheet's 'Fitting' (SGL (Area)), with a Smart background
    over the 10-90 % span of its B.E. range."""
    if not peak_fitting_by_scan:
        return
    for sheet_name, cl in levels.items():
        m = re.match(r'([A-Za-z]+\d*[spdfg]*)', sheet_name)
        base = m.group(1) if m else sheet_name
        if base not in peak_fitting_by_scan:
            continue
        fitting = cl.setdefault('Fitting', {})
        fitting['Peaks'] = {}
        fitting['Fitting Model'] = 'SGL (Area)'
        for peak in peak_fitting_by_scan[base]:
            pos_c = peak.get('position_constraint') or \
                f"{peak['position'] - 2:.2f}:{peak['position'] + 2:.2f}"
            fwhm_c = peak.get('fwhm_constraint') or "0.30:3.50"
            lg_c = peak.get('lg_constraint') or "0.00:100.00"
            fitting['Peaks'][peak['name']] = {
                'Position': float(f"{peak['position']:.2f}"),
                'Height': float(f"{peak['height']:.2f}"),
                'FWHM': float(f"{peak['fwhm']:.2f}"),
                'L/G': float(f"{peak['lg_ratio']:.2f}"),
                'Area': float(f"{peak['area']:.2f}"),
                'Sigma': 0.00,
                'Gamma': 0.00,
                'Skew': 0.00,
                'Fitting Model': 'SGL (Area)',
                'Constraints': {
                    'Position': pos_c,
                    'Height': f"{max(0, peak['height'] * 0.5):.2f}:"
                              f"{peak['height'] * 1.5:.2f}",
                    'FWHM': fwhm_c,
                    'L/G': lg_c,
                    'Area': f"{max(0, peak['area'] * 0.5):.2f}:"
                            f"{peak['area'] * 1.5:.2f}",
                    'Sigma': "0.10:1.00",
                    'Gamma': "0.10:1.00",
                    'Skew': "0.01:2.00",
                },
            }

        x = np.array(cl['B.E.'])
        y = np.array(cl['Raw Data'])
        try:
            x_min, x_max = np.min(x), np.max(x)
        except ValueError:
            continue    # KherveFitting would fail on an empty sheet too
        span = x_max - x_min
        lo = float(f"{(x_min + 0.1 * span):.2f}")
        hi = float(f"{(x_min + 0.9 * span):.2f}")
        if lo > hi:
            lo, hi = hi, lo
        try:
            bkg = _adaptive_smart_background(x, y, (lo, hi), y.copy(),
                                             0.0, 0.0)
            background = cl.setdefault('Background', {})
            background['Bkg Y'] = bkg.tolist()
            background['Bkg Low'] = lo
            background['Bkg High'] = hi
            background['Recorded_Ranges'] = [(0.00, 0.00, lo, hi)]
        except Exception:
            cl['Background'] = {'Bkg Y': y.tolist(), 'Bkg Low': lo,
                                'Bkg High': hi,
                                'Recorded_Ranges': [(0.00, 0.00, lo, hi)]}


# --- building the intermediate sheets --------------------------------------

def _exp_rows(title_info, acquisition_params):
    return list(title_info.items()) + list(acquisition_params.items())


def _avantage_xlsx_sheets(wb):
    """``import_avantage_file_direct``: the '_Kfitting.xlsx' sheets it
    writes, as [(title, rows, exp_rows)]."""
    groups = {}
    for name in wb.sheetnames:
        if _is_spectrum_sheet(name, _SPECTRUM_TERMS_XLSX):
            parts = name.split()
            suffix = (parts[-1] if len(parts) >= 2 and len(parts[-1]) == 1
                      and parts[-1].isalpha() else None)
            groups.setdefault(suffix, []).append(name)

    out = []
    titles = []

    def new_sheet(title, rows, exp):
        title = _xl_title(titles, title)
        titles.append(title)
        out.append((title, rows, exp))

    sample_counter = 0
    for suffix in sorted(groups, key=lambda s: (s is not None, s)):
        title_info = extract_title_sheet_info(wb, suffix)
        vb = fermi = cutoff = 0
        for sheet_name in groups[suffix]:
            sheet = wb[sheet_name]
            acq = extract_acquisition_parameters(sheet)
            lower = sheet_name.lower()
            if "Survey" in sheet_name or "survey" in sheet_name:
                base = "Survey"
            elif any(t in lower for t in _VB_TERMS):
                base = "VB" if vb == 0 else f"VB{vb + 1}"
                vb += 1
            elif any(t in lower for t in ['fermi', 'fermi scan']):
                base = "Fermi" if fermi == 0 else f"Fermi{fermi + 1}"
                fermi += 1
            elif any(t in lower for t in ['cut-off', 'cutoff', 'cut off']):
                base = "Cut-Off" if cutoff == 0 else f"Cut-Off{cutoff + 1}"
                cutoff += 1
            else:
                base = sheet_name.split()[0]
            final = base if sample_counter == 0 else f"{base}{sample_counter}"

            num_samples, multi = 1, False
            d8 = sheet.cell(row=8, column=4).value
            if d8 is not None:
                try:
                    num_samples = int(d8)
                    multi = num_samples > 1
                except (ValueError, TypeError):
                    pass

            start_row = 19
            for r in range(1, sheet.max_row + 1):
                if sheet.cell(row=r, column=1).value == "eV":
                    start_row = r + 1
                    break

            exp = (_exp_rows(title_info, acq)
                   if (acq or title_info) else [])
            if multi:
                for sidx in range(num_samples):
                    if sample_counter == 0:
                        title = f"{base}{sidx if sidx > 0 else ''}"
                    elif sidx == 0:
                        title = f"{base}{sample_counter}"
                    else:
                        title = f"{base}{sample_counter}_{sidx}"
                    rows = []
                    for r in range(start_row, sheet.max_row + 1):
                        be = sheet.cell(row=r, column=1).value
                        iv = sheet.cell(row=r, column=3 + sidx).value
                        try:
                            float(be)
                            float(iv)
                        except (ValueError, TypeError):
                            continue
                        rows.append((f"{float(be):.2f}", f"{iv}"))
                    new_sheet(title, rows, exp)
            else:
                title = final
                num = re.search(r'\((\d+)\)', sheet_name)
                if num:
                    title = (f"{base}{num.group(1)}" if sample_counter == 0
                             else f"{base}{sample_counter}_{num.group(1)}")
                rows = []
                for r in range(start_row, sheet.max_row + 1):
                    be = sheet.cell(row=r, column=1).value
                    iv = sheet.cell(row=r, column=3).value
                    try:
                        float(be)
                        float(iv)
                    except (ValueError, TypeError):
                        continue
                    rows.append((f"{float(be):.2f}", f"{float(iv):.2f}"))
                new_sheet(title, rows, exp)
        sample_counter += 1
    return out


def _avantage_xls_sheets(wb_xls):
    """``import_avantage_file_direct_xls``."""
    groups = {}
    for name in wb_xls.sheet_names():
        if _is_spectrum_sheet(name, _SPECTRUM_TERMS_XLS):
            parts = name.split()
            suffix = (parts[-1] if len(parts) >= 2 and len(parts[-1]) == 1
                      and parts[-1].isalpha() else None)
            groups.setdefault(suffix, []).append(name)

    out, titles = [], []
    sample_counter = 0
    for suffix in sorted(groups, key=lambda s: (s is not None, s)):
        title_info = extract_title_sheet_info_xls(wb_xls, suffix)
        vb = fermi = cutoff = 0
        for sheet_name in groups[suffix]:
            sheet = wb_xls.sheet_by_name(sheet_name)
            acq = extract_acquisition_parameters(_TempSheet(sheet))
            lower = sheet_name.lower()
            num = re.search(r'\((\d+)\)', sheet_name)
            if "Survey" in sheet_name or "survey" in sheet_name:
                if num:
                    title = (f"Survey{num.group(1)}" if sample_counter == 0
                             else f"Survey{sample_counter}_{num.group(1)}")
                else:
                    title = ("Survey" if sample_counter == 0
                             else f"Survey{sample_counter}")
            elif any(t in lower for t in _VB_TERMS):
                base = "VB" if vb == 0 else f"VB{vb + 1}"
                vb += 1
                title = base if sample_counter == 0 \
                    else f"{base}{sample_counter}"
            elif any(t in lower for t in ['fermi', 'fermi scan']):
                base = "Fermi" if fermi == 0 else f"Fermi{fermi + 1}"
                fermi += 1
                title = base if sample_counter == 0 \
                    else f"{base}{sample_counter}"
            elif any(t in lower for t in ['cut-off', 'cutoff', 'cut off']):
                base = "Cut-Off" if cutoff == 0 else f"Cut-Off{cutoff + 1}"
                cutoff += 1
                title = base if sample_counter == 0 \
                    else f"{base}{sample_counter}"
            else:
                element = sheet_name.split()[0]
                if num:
                    title = (f"{element}{num.group(1)}" if sample_counter == 0
                             else f"{element}{sample_counter}_{num.group(1)}")
                else:
                    title = (element if sample_counter == 0
                             else f"{element}{sample_counter}")

            start_row = 17
            for r in range(sheet.nrows):
                if sheet.cell_value(r, 0) == "eV":
                    start_row = r + 1
                    break
            rows = []
            for r in range(start_row, sheet.nrows):
                values = sheet.row_values(r)
                if len(values) < 3:
                    continue
                try:
                    be, iv = float(values[0]), float(values[2])
                except (ValueError, TypeError):
                    continue
                rows.append((f"{be:.2f}", f"{iv:.2f}"))
            exp = _exp_rows(title_info, acq) if acq else []
            title = _xl_title(titles, title)
            titles.append(title)
            out.append((title, rows, exp))
        sample_counter += 1
    return out


def _load_sheets(sheets):
    """``open_xlsx_file`` reading the intermediate workbook back."""
    levels = {}
    for title, rows, exp in sheets:
        levels[title] = _sheet_core_level(title, rows, exp)
    return levels


def _no_spectra(path):
    return ValueError(f"{os.path.basename(str(path))} has no Avantage "
                      f"spectrum sheets ('Survey', '… Scan', VB, …).")


def read_avantage(path):
    """An Avantage export (.xlsx with peaks, or .xls) -> its spectra."""
    path = str(path)
    if path.lower().endswith('.xls'):
        import xlrd
        wb_xls = xlrd.open_workbook(path)
        sheets = _avantage_xls_sheets(wb_xls)
        if not sheets:
            raise _no_spectra(path)
        return _load_sheets(sheets)

    import openpyxl
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:
        raise ValueError(f"{os.path.basename(path)} is not an Excel "
                         f"workbook: {exc}") from exc
    try:
        sheets = _avantage_xlsx_sheets(wb)
        peaks = extract_peak_fitting_from_peak_table(wb)
    finally:
        wb.close()
    if not sheets:
        raise _no_spectra(path)
    levels = _load_sheets(sheets)
    add_peak_fitting(levels, peaks)
    return levels


def _multi_xlsx_sheets(wb, sample_idx, titles):
    """``process_avantage_xlsx_with_sample_number``."""
    out = []
    for sheet_name in wb.sheetnames:
        if not _is_spectrum_sheet(sheet_name, _SPECTRUM_TERMS_XLSX):
            continue
        sheet = wb[sheet_name]
        acq = extract_acquisition_parameters(sheet)
        base = ("Survey" if "Survey" in sheet_name or "survey" in sheet_name
                else sheet_name.split()[0])
        num_samples, multi = 1, False
        d8 = sheet.cell(row=8, column=4).value
        if d8 is not None:
            try:
                num_samples = int(d8)
                multi = num_samples > 1
            except (ValueError, TypeError):
                pass
        start_row = 19
        for r in range(1, sheet.max_row + 1):
            if sheet.cell(row=r, column=1).value == "eV":
                start_row = r + 1
                break
        exp = list(acq.items())
        for sub in range(num_samples if multi else 1):
            title = f"{base}{sample_idx}"
            if multi and sub > 0:
                title = f"{base}{sample_idx}_{sub}"
            rows = []
            for r in range(start_row, sheet.max_row + 1):
                be = sheet.cell(row=r, column=1).value
                iv = sheet.cell(row=r, column=3 + sub).value
                if be is None or iv is None:
                    continue
                rows.append((be, iv))
            title = _xl_title(titles, title)
            titles.append(title)
            out.append((title, rows, exp))
    return out


def _multi_xls_sheets(wb_xls, sample_idx, titles):
    """``process_avantage_xls_with_sample_number`` (data from row 19)."""
    out = []
    for sheet_name in wb_xls.sheet_names():
        if not _is_spectrum_sheet(sheet_name, _SPECTRUM_TERMS_XLSX):
            continue
        sheet = wb_xls.sheet_by_name(sheet_name)
        acq = extract_acquisition_parameters(_TempSheet(sheet))
        base = ("Survey" if "Survey" in sheet_name or "survey" in sheet_name
                else sheet_name.split()[0])
        rows = []
        for r in range(18, sheet.nrows):
            be = sheet.cell_value(r, 0)
            iv = sheet.cell_value(r, 2) if sheet.ncols > 2 else ''
            if be and iv:
                try:
                    rows.append((f"{be:.2f}", f"{iv:.2f}"))
                except (ValueError, TypeError):
                    continue   # KherveFitting stops on a text cell here
        title = _xl_title(titles, f"{base}{sample_idx}")
        titles.append(title)
        out.append((title, rows, list(acq.items())))
    return out


def read_avantage_files(paths):
    """Several Avantage exports, one sample each (KherveFitting 'Avantage
    Multiple xlsx files (folder)'): files in name order, sheets numbered by
    sample — C1s0, O1s0, C1s1, ... No peaks, no Titles info (as there)."""
    ordered = sorted((str(p) for p in paths), key=os.path.basename)
    sheets, titles = [], []
    for idx, path in enumerate(ordered):
        if path.lower().endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(path)
            try:
                sheets += _multi_xlsx_sheets(wb, idx, titles)
            finally:
                wb.close()
        else:
            import xlrd
            sheets += _multi_xls_sheets(xlrd.open_workbook(path), idx,
                                        titles)
    if not sheets:
        raise _no_spectra(ordered[0] if ordered else "")
    return _load_sheets(sheets)


def sniff_avantage(path):
    """An Excel workbook with a 'Titles' sheet (Avantage's export)."""
    path = str(path)
    if path.lower().endswith('.xls'):
        try:
            import xlrd
        except ImportError:
            return False
        try:
            return "Titles" in xlrd.open_workbook(
                path, on_demand=True).sheet_names()
        except Exception:
            return False
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:
        return False
    try:
        return "Titles" in wb.sheetnames
    finally:
        wb.close()
