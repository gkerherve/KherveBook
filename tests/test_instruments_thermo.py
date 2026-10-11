"""Thermo XPS readers (khervebook.instruments.thermo): VGD, AVG, Avantage.

Real files are used when they are on this machine (skipped otherwise);
the multi-spectrum / area-scan VGD, AVG and .xls cases use small
synthetic files built here in the layout KherveFitting's parsers read.

Copyright (C) 2026 Gwilherm Kerherve

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.
"""

import glob
import os
import struct

import numpy as np
import pytest

from khervebook import kfitio
from khervebook.instruments import kfitwrite, thermo

WS = "/Users/gkerherv/Desktop/WS_fits"
SIGMA = os.path.join(WS, "Sigmaaldrich", "Sigmaaldrich_c.DATA")
AVANTAGE_DIR = ("/Users/gkerherv/Documents/PycharmProjects/KherveFittingPro/"
                "Data-Examples/zz RawData_ToBeImported")


def _real(path):
    if not os.path.exists(path):
        pytest.skip(f"real file not present: {path}")
    return path


def _roundtrip(levels):
    project = kfitio.read_bytes(kfitwrite.to_bytes(levels, "x.kfit"))
    assert project.names == list(levels)
    return project


# ---------------------------------------------------------------------------
# synthetic file builders (also used by the out-of-tree parity check)
# ---------------------------------------------------------------------------

_END, _FREE, _FATSECT, _NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, \
    0xFFFFFFFF


def write_cfb(path, streams):
    """A minimal OLE2 compound file holding *streams* {name: bytes} in the
    root storage (version 3, 512-byte sectors, real mini stream)."""
    names = list(streams)
    sectors = []        # list of 512-byte blocks
    fat = []

    def put_chain(data):
        if not data:
            return _END
        n = (len(data) + 511) // 512
        first = len(sectors)
        for i in range(n):
            sectors.append(data[i * 512:(i + 1) * 512].ljust(512, b"\0"))
            fat.append(first + i + 1 if i < n - 1 else _END)
        return first

    mini, minifat, starts = b"", [], {}
    for name in names:
        data = streams[name]
        if len(data) < 4096:
            n = (len(data) + 63) // 64
            first = len(mini) // 64
            mini += data.ljust(n * 64, b"\0")
            minifat += [first + i + 1 if i < n - 1 else _END
                        for i in range(n)]
            starts[name] = first if n else _END
        else:
            starts[name] = put_chain(data)
    root_start = put_chain(mini)
    minifat_bytes = b"".join(struct.pack("<I", v) for v in minifat)
    minifat_start = put_chain(minifat_bytes) if minifat else _END

    def entry(name, etype, child, right, start, size):
        raw = name.encode("utf-16-le") + b"\0\0"
        e = raw.ljust(64, b"\0")
        e += struct.pack("<HBB", len(raw), etype, 1)
        e += struct.pack("<III", _NOSTREAM, right, child)
        e += b"\0" * 16 + b"\0" * 4 + b"\0" * 16
        e += struct.pack("<IQ", start, size)
        return e

    entries = [entry("Root Entry", 5, 1 if names else _NOSTREAM, _NOSTREAM,
                     root_start, len(mini))]
    for i, name in enumerate(names):
        right = i + 2 if i + 1 < len(names) else _NOSTREAM
        entries.append(entry(name, 2, _NOSTREAM, right, starts[name],
                             len(streams[name])))
    while len(entries) % 4:
        entries.append(b"\0" * 64 + struct.pack("<HBB", 0, 0, 0)
                       + struct.pack("<III", _NOSTREAM, _NOSTREAM, _NOSTREAM)
                       + b"\0" * 36 + struct.pack("<IQ", 0, 0))
    dir_start = put_chain(b"".join(entries))

    # FAT sectors go last; they describe themselves too.
    n_fat = 1
    while (len(sectors) + n_fat) > n_fat * 128:
        n_fat += 1
    fat_first = len(sectors)
    fat += [_FATSECT] * n_fat
    fat += [_FREE] * (n_fat * 128 - len(fat))
    fat_bytes = b"".join(struct.pack("<I", v) for v in fat)
    for i in range(n_fat):
        sectors.append(fat_bytes[i * 512:(i + 1) * 512])

    header = bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 16
    header += struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6) + b"\0" * 6
    header += struct.pack("<IIIIIIIII", 0, n_fat, dir_start, 0, 4096,
                          minifat_start, (len(minifat) * 4 + 511) // 512,
                          _END, 0)
    difat = [fat_first + i for i in range(n_fat)] + [_FREE] * (109 - n_fat)
    header += b"".join(struct.pack("<I", v) for v in difat)
    with open(path, "wb") as fh:
        fh.write(header + b"".join(sectors))


def vgd_bytes(intensities, dims, ke_start=1180.0, ke_step=0.5,
              xy=(0.0, 20.0, 0.0, 30.0), pass_energy=20.0, dwell=0.05,
              periods=10):
    """The four streams parse_vgd_file reads. *dims* = (n_ke,) for one
    spectrum, (n_ke, n_spectra) for 2-D, (n_ke, n_x, n_y) for an area."""
    data = np.asarray(intensities, dtype="<f8").tobytes()
    axes = bytearray(56 if len(dims) == 3 else 32)
    struct.pack_into("<i", axes, 4, len(dims))
    for k, d in enumerate(dims):
        struct.pack_into("<i", axes, 12 + 16 * k, d - 1)
    space = bytearray(140)
    struct.pack_into("<dd", space, 30, ke_start, ke_step)
    struct.pack_into("<dd", space, 73, xy[0], xy[1])
    struct.pack_into("<dd", space, 116, xy[2], xy[3])
    prop = bytearray(4000)
    struct.pack_into("<f", prop, 100, 1486.6)
    struct.pack_into("<i", prop, 168, periods)
    struct.pack_into("<f", prop, 200, pass_energy)
    struct.pack_into("<f", prop, 208, dwell)
    struct.pack_into("<f", prop, 216, 4.2)
    for j, c in enumerate((4.1, 0.73, -0.67, 0.08)):
        struct.pack_into("<f", prop, 3200 + 8 * j, c)
    return {"VGData": data, "VGDataAxes": bytes(axes),
            "VGSpaceAxes": bytes(space), thermo._PROP_STREAM: bytes(prop)}


def write_vgd(path, intensities, dims, **kw):
    write_cfb(path, vgd_bytes(intensities, dims, **kw))


AVG_TEXT = """\
$PROPERTIES=
DS_EXT_SUPROPID_TITLE           : VT_BSTR = 'Ti2p Snap'
DS_EXT_SUPROPID_SUBJECT         : VT_BSTR = 'Sample42'
DS_EXT_SUPROPID_AUTHOR          : VT_BSTR = 'tester'
DS_EXT_SUPROPID_CREATED         : VT_DATE = 01/02/2024 10:11:12
DS_EXT_SUPROPID_SAVED           : VT_DATE = 01/02/2024 11:12:13
DS_GEPROPID_INSTRUMENT          : VT_BSTR = 'K-Alpha'
DS_SOPROPID_ENERGY              : VT_R4 = 1486.68
DS_ANPROPID_PASS                : VT_R4 = 50
DS_ANPROPID_WORK_FTN            : VT_R4 = 4.3
DS_ACPROPID_ACQ_TIME            : VT_R4 = 0.05
DS_ACPROPID_PERIODS             : VT_I4 = 4
DS_ANPROPID_LENS_MODE_NAME      : VT_BSTR = 'Standard'
DS_ANPROPID_TXFN_COEFF[0]       : VT_R4 = 4.1
DS_ANPROPID_TXFN_COEFF[1]       : VT_R4 = 0.7
$DATAAXES=2,#empty#
$SPACEAXES=2
   0=   1000.000000,   0.100000,   6,   ENERGY,   LINEAR,   'E',   'eV',   'Kinetic Energy'
   1=   0.000000,   30.000000,   2,   ETCHTIME,   LINEAR,   't',   's',   'Etch Time'
$AXISVALUE= DATAXIS=1 SPACEAXIS=1 LABEL='Etch Time' POINT=0 VALUE=0.000000;
$AXISVALUE= DATAXIS=1 SPACEAXIS=1 LABEL='Etch Time' POINT=1 VALUE=30.000000;
$DATA=*
LIST@    0=  100.5,  110.25,  120.125,  130.0
LIST@    4=  140.0,  150.0
$DATA=*
LIST@    0=  200.0,  210.0,  220.0,  230.0
LIST@    4=  240.0
"""


def write_avg(path, text=AVG_TEXT):
    with open(path, "w", encoding="latin-1") as fh:
        fh.write(text)


def write_avantage_xlsx(path, peaks=True):
    """An Avantage-style export: Titles, two spectra, a Peak Table."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, be0, n in (("Survey", 1100.0, 30), ("C1s Scan", 295.0, 50)):
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=8, value="Acquisition Parameters :")
        ws.cell(row=2, column=8, value="Parameter")
        ws.cell(row=3, column=8, value="Number of Scans")
        ws.cell(row=3, column=9, value=5)
        ws.cell(row=4, column=8, value="Lens Mode")
        ws.cell(row=4, column=9, value="Standard")
        ws.cell(row=18, column=1, value="eV")
        ws.cell(row=18, column=3, value="Counts / s")
        x = be0 - 0.2 * np.arange(n)
        y = 1000 + 5000 * np.exp(-((x - (be0 - 4)) / 0.8) ** 2) \
            + 20 * np.arange(n)[::-1]
        for i, (a, b) in enumerate(zip(x, y)):
            ws.cell(row=19 + i, column=1, value=round(float(a), 2))
            ws.cell(row=19 + i, column=3, value=round(float(b), 3))
    t = wb.create_sheet("Titles")
    t["A1"] = "My sample"
    t["A3"] = "Grid 1"
    if peaks:
        pt = wb.create_sheet("Peak Table")
        pt["A1"] = "Peak Fit Table : C1s Scan"
        pt["A2"] = "Ref"
        rows = [("L", "C1s C-C", 291.0, 4000.0, None, 6000.0, None, 1.2, 30),
                (None, None, "280:300:00", None, None, None, None,
                 "0.5:2:00", None),
                ("M", "C1s C-O", 289.5, 1000.0, None, 1500.0, None, 1.3, 30),
                (None, None, "L+1.5", None, None, None, None, "L*1", None)]
        for r, row in enumerate(rows, start=3):
            for c, v in enumerate(row, start=1):
                if v is not None:
                    pt.cell(row=r, column=c, value=v)
    wb.save(path)


def write_avantage_xls(path):
    import xlwt
    wb = xlwt.Workbook()
    for name, be0 in (("XPS Survey", 1100.0), ("O1s Scan (2)", 540.0)):
        ws = wb.add_sheet(name)
        ws.write(0, 7, "Acquisition Parameters :")
        ws.write(2, 7, "Number of Scans")
        ws.write(2, 8, 3)
        ws.write(17, 0, "eV")
        for i in range(12):
            ws.write(18 + i, 0, be0 - 0.5 * i)
            ws.write(18 + i, 2, 100.0 + i * 1.234)
    wb.add_sheet("Titles").write(0, 0, "XLS sample")
    wb.save(path)


# ---------------------------------------------------------------------------
# registry / sniffing
# ---------------------------------------------------------------------------

def test_formats():
    by_label = {f.label: f for f in thermo.FORMATS}
    assert set(by_label) == {"Avantage (.xlsx, .xls)",
                             "Avantage files (.xlsx, .xls)", "VGD (.vgd)",
                             "VGD files (.vgd)", "AVG (.avg)",
                             "AVG files (.avg)"}
    for f in thermo.FORMATS:
        assert f.technique == "XPS" and f.group == "Thermo"
    assert by_label["VGD (.vgd)"].needs == ("olefile",)
    assert by_label["VGD files (.vgd)"].multi
    assert by_label["VGD files (.vgd)"].sniff("x.vgd") is False


def test_sniff_avantage(tmp_path):
    p = tmp_path / "a.xlsx"
    write_avantage_xlsx(p)
    assert thermo.sniff_avantage(p)
    import openpyxl
    q = tmp_path / "plain.xlsx"
    openpyxl.Workbook().save(q)
    assert not thermo.sniff_avantage(q)
    c = tmp_path / "x.csv"
    c.write_text("1,2\n")
    assert not thermo.sniff_avantage(c)


# ---------------------------------------------------------------------------
# VGD
# ---------------------------------------------------------------------------

def test_vgd_real_o1s():
    path = _real(os.path.join(SIGMA, "O1s Scan.VGD"))
    levels = thermo.read_vgd(path)
    assert list(levels) == ["O1s"]
    cl = levels["O1s"]
    assert len(cl["B.E."]) == 151
    assert cl["B.E."][0] > cl["B.E."][-1]          # high -> low BE
    assert cl["B.E."][0] == pytest.approx(540.08, abs=0.01)
    assert cl["B.E."][-1] == pytest.approx(525.08, abs=0.01)
    # Raw Data = counts / (periods x dwell); Corrected Data = counts
    factor = cl["Transmission"][0]
    assert factor == pytest.approx(10 * 0.05)
    assert cl["Raw Data"][5] == pytest.approx(cl["Corrected Data"][5]
                                              / factor)
    info = cl["ExperimentalInfo"]
    assert info["Species & Transition"] == "O1s"
    assert info["Pass Energy"] == "20.00"
    assert info["TXF Applied"] == "Yes"
    assert info["Title"] == "O1s Scan"
    project = _roundtrip(levels)
    assert np.allclose(project.sheet("O1s").x, cl["B.E."])


def test_vgd_real_survey_name():
    path = _real(os.path.join(SIGMA, "XPS Survey.VGD"))
    levels = thermo.read_vgd(path)
    assert list(levels) == ["XPS Survey"]
    assert len(levels["XPS Survey"]["B.E."]) == 2721


def test_vgd_real_folder():
    files = sorted(glob.glob(os.path.join(SIGMA, "*.VGD")))
    if not files:
        pytest.skip("real VGD folder not present")
    levels = thermo.read_vgd_files(files + files[:1])
    names = list(levels)
    first = thermo.extract_vgd_core_level_name(os.path.basename(files[0]))
    assert names[0] == first and names[-1] == first + "1"
    assert len(names) == len(files) + 1
    _roundtrip(levels)


def test_vgd_multi_spectrum(tmp_path):
    n_ke, n_s = 20, 3
    data = np.concatenate([np.full(n_ke, 100.0 * (s + 1)) + np.arange(n_ke)
                           for s in range(n_s)])
    p = tmp_path / "Pt4f Scan.vgd"
    write_vgd(p, data, (n_ke, n_s))
    levels = thermo.read_vgd(p)
    assert list(levels) == ["Pt4f", "Pt4f1", "Pt4f2"]
    cl = levels["Pt4f2"]
    assert len(cl["B.E."]) == n_ke
    se = struct.unpack("<f", struct.pack("<f", 1486.6))[0]
    assert cl["B.E."][0] == pytest.approx(se - 1180.0)
    assert cl["B.E."][1] - cl["B.E."][0] == pytest.approx(-0.5)
    assert cl["Corrected Data"][0] == 300.0
    assert cl["Raw Data"][0] == pytest.approx(300.0 / (10 * 0.05))
    info = cl["ExperimentalInfo"]
    assert info["Spectrum Index"] == "2" and info["Total Spectra"] == "3"
    assert info["Title"] == "Unknown" and info["Sample ID"] == "Pt4f Scan"
    assert info["Periods"] == "10" and info["Work Function"] == "4.20"
    assert info["Source Label"] == "Al K-alpha Monochromated"
    _roundtrip(levels)


def test_vgd_area_scan(tmp_path):
    n_ke, n_x, n_y = 10, 3, 2
    arr = np.zeros((n_y, n_x, n_ke))
    for y in range(n_y):
        for x in range(n_x):
            arr[y, x] = 10 * (y * n_x + x) + np.arange(n_ke)
    p = tmp_path / "C1s Scan.vgd"
    write_vgd(p, arr.ravel(), (n_ke, n_x, n_y))
    levels = thermo.read_vgd(p)
    assert list(levels) == ["C1s~Map", "C1s"]
    m = levels["C1s~Map"]
    assert m["_Map_type"] == "vgd_xy" and m["_num_sweeps"] == 6
    assert (m["_n_x"], m["_n_y"]) == (3, 2)
    assert m["_x_step"] == 20.0 and m["_y_step"] == 30.0
    assert m["Y5"][0] == pytest.approx(40 / 0.5)        # pixel y=1, x=1
    assert m["ExperimentalInfo"]["Map Kind"] == "XY"
    avg = levels["C1s"]
    assert avg["Raw Data"][0] == pytest.approx(25 / 0.5)
    assert avg["ExperimentalInfo"]["Source Map"] == "C1s~Map"
    assert "Map Kind" not in avg["ExperimentalInfo"]
    _roundtrip(levels)


def test_vgd_rejects_other_files(tmp_path):
    p = tmp_path / "bad.vgd"
    p.write_bytes(b"not an ole file at all")
    with pytest.raises(ValueError):
        thermo.read_vgd(p)
    q = tmp_path / "empty.vgd"
    write_cfb(q, {"Other": b"x" * 10})
    with pytest.raises(ValueError):
        thermo.read_vgd(q)


# ---------------------------------------------------------------------------
# AVG
# ---------------------------------------------------------------------------

def test_avg_depth_profile(tmp_path):
    p = tmp_path / "Ti2p Snap.avg"
    write_avg(p)
    levels = thermo.read_avg(p)
    assert list(levels) == ["Ti2p", "Ti2p1"]
    cl = levels["Ti2p"]
    assert cl["B.E."] == [486.68, 486.58, 486.48, 486.38, 486.28, 486.18]
    assert cl["Corrected Data"] == [100.5, 110.25, 120.12, 130.0, 140.0,
                                    150.0]
    assert cl["Raw Data"][0] == pytest.approx(round(100.5 / 0.2, 2))
    assert cl["Transmission"] == [0.2] * 6
    second = levels["Ti2p1"]
    assert second["Corrected Data"][-1] == 0.0          # padded to 6 points
    info = second["ExperimentalInfo"]
    assert info["Etch Time (s)"] == "30.0000"
    assert info["Sample ID"] == "Sample42"
    assert info["Date Created"] == "01/02/2024"
    assert info["TXF Coefficients"] == "4.100000, 0.700000"
    assert info["Lens Mode"] == "Standard"
    assert cl["Background"]["Bkg Low"] == 486.18
    _roundtrip(levels)


def test_avg_files_naming(tmp_path):
    p = tmp_path / "Ti2p Snap.avg"
    q = tmp_path / "Ti2p_Scan.avg"
    write_avg(p)
    write_avg(q)
    levels = thermo.read_avg_files([p, q])
    assert list(levels) == ["Ti2p", "Ti2p1", "Ti2p2", "Ti2p3"]


def test_avg_rejects_text(tmp_path):
    p = tmp_path / "x.avg"
    p.write_text("hello\nworld\n")
    with pytest.raises(ValueError):
        thermo.read_avg(p)


# ---------------------------------------------------------------------------
# Avantage
# ---------------------------------------------------------------------------

def test_avantage_real():
    files = sorted(glob.glob(os.path.join(AVANTAGE_DIR, "*_Avantage.xlsx")))
    if not files:
        pytest.skip("real Avantage workbooks not present")
    path = os.path.join(AVANTAGE_DIR, "Au_Metal_Avantage.xlsx")
    if path not in files:
        path = files[0]
    assert thermo.sniff_avantage(path)
    levels = thermo.read_avantage(path)
    assert list(levels)[0] == "Survey"
    survey = levels["Survey"]
    assert len(survey["B.E."]) == 2721
    assert survey["B.E."][0] == 1350.08 and survey["B.E."][1] == 1349.58
    assert survey["Corrected Data"] == survey["Raw Data"]
    assert set(survey["Transmission"]) == {1.0}
    info = survey["ExperimentalInfo"]
    assert info["Lens Mode"] == "Standard"
    assert "Sample Info" in info
    _roundtrip(levels)


def test_avantage_peaks(tmp_path):
    p = tmp_path / "s_Avantage.xlsx"
    write_avantage_xlsx(p)
    levels = thermo.read_avantage(p)
    assert list(levels) == ["Survey", "C1s"]
    c1s = levels["C1s"]
    assert len(c1s["B.E."]) == 50 and c1s["B.E."][0] == 295.0
    info = c1s["ExperimentalInfo"]
    assert info == {"Sample Info": "My sample", "Additional Info": "Grid 1",
                    "Number of Scans": "5", "Lens Mode": "Standard"}
    peaks = c1s["Fitting"]["Peaks"]
    assert list(peaks) == ["C1s C-C", "C1s C-O"]
    assert peaks["C1s C-C"]["Constraints"]["Position"] == "280:300"
    assert peaks["C1s C-O"]["Constraints"]["Position"] == "A+1.5"
    assert peaks["C1s C-O"]["Constraints"]["FWHM"] == "A*1"
    assert peaks["C1s C-O"]["Constraints"]["L/G"] == "0.00:100.00"
    assert peaks["C1s C-C"]["Area"] == 6000.0
    bkg = c1s["Background"]
    assert bkg["Bkg Low"] == pytest.approx(286.18)
    assert bkg["Bkg High"] == pytest.approx(294.02)
    assert len(bkg["Bkg Y"]) == 50
    assert "Fitting" not in levels["Survey"]
    project = _roundtrip(levels)
    assert [pk["name"] for pk in project.sheet("C1s").peaks] == \
        ["C1s C-C", "C1s C-O"]


def test_avantage_files(tmp_path):
    a, b = tmp_path / "b.xlsx", tmp_path / "a.xlsx"
    write_avantage_xlsx(a, peaks=False)
    write_avantage_xlsx(b)
    levels = thermo.read_avantage_files([a, b])
    assert list(levels) == ["Survey0", "C1s0", "Survey1", "C1s1"]
    assert "Fitting" not in levels["C1s0"]
    assert "Sample Info" not in levels["C1s0"]["ExperimentalInfo"]
    _roundtrip(levels)


def test_avantage_xls(tmp_path):
    pytest.importorskip("xlrd")
    pytest.importorskip("xlwt")
    p = tmp_path / "s.xls"
    write_avantage_xls(p)
    assert thermo.sniff_avantage(p)
    levels = thermo.read_avantage(p)
    assert list(levels) == ["Survey", "O1s2"]
    o1s = levels["O1s2"]
    assert o1s["B.E."][:2] == [540.0, 539.5]
    assert o1s["Raw Data"][1] == 101.23
    assert o1s["ExperimentalInfo"] == {"Sample Info": "XLS sample",
                                       "Number of Scans": "3.0"}
    _roundtrip(levels)


def test_avantage_rejects_plain_workbook(tmp_path):
    import openpyxl
    p = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(p)
    with pytest.raises(ValueError):
        thermo.read_avantage(p)
