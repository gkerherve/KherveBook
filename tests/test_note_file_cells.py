"""Note (rich text + pen) and File (attachment) cell tests.

Copyright (C) 2026 Gwilherm Kerherve

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
(at your option) any later version.
"""

import json
from pathlib import Path


# -- Note cell -------------------------------------------------------------

def test_note_cell_html_round_trip(qapp):
    from khervebook.notecell import NoteCell
    cell = NoteCell()
    cell.set_source(json.dumps({"kbook_note": 1,
                                "html": "<h2>Hi</h2><p>body</p>",
                                "ink": None}))
    src = cell.source()
    doc = json.loads(src)
    assert doc["kbook_note"] == 1
    assert "Hi" in doc["html"] and "body" in doc["html"]


def test_note_cell_ink_round_trip(qapp):
    from khervebook.notecell import NoteCell
    ink = {"ref_w": 400.0,
           "strokes": [{"color": "#c0392b", "width": 3.0,
                        "pts": [[1.0, 2.0], [3.0, 4.0]]}]}
    cell = NoteCell()
    cell.set_source(json.dumps({"kbook_note": 1, "html": "<p>x</p>",
                                "ink": ink}))
    doc = json.loads(cell.source())
    assert doc["ink"]["ref_w"] == 400.0
    assert len(doc["ink"]["strokes"]) == 1
    assert doc["ink"]["strokes"][0]["pts"] == [[1.0, 2.0], [3.0, 4.0]]


def test_note_cell_formatting_toggles_bold(qapp):
    from khervebook.notecell import NoteCell
    from PyQt5.QtGui import QTextCursor, QFont
    cell = NoteCell()
    cell.rich.setPlainText("hello world")
    cur = cell.rich.textCursor()
    cur.select(QTextCursor.Document)
    cell.rich.setTextCursor(cur)
    cell.toggle_bold()
    assert cell.rich.textCursor().charFormat().fontWeight() == QFont.Bold


def test_note_cell_registered_and_round_trips_in_notebook(qapp):
    from khervebook.notebook import NotebookWidget
    nb = NotebookWidget()
    note = nb.add_cell("note")
    note.set_source(json.dumps({"kbook_note": 1, "html": "<p>keep me</p>",
                                "ink": None}))
    reloaded = NotebookWidget()
    reloaded.load_json(nb.to_json())
    got = [c for c in reloaded.cells if c.CELL_TYPE == "note"][0]
    assert "keep me" in json.loads(got.source())["html"]


# -- File cell -------------------------------------------------------------

def test_file_cell_small_file_embeds(qapp, tmp_path):
    from khervebook.filecell import FileCell
    f = tmp_path / "data.csv"
    f.write_text("a,b\n1,2\n")
    cell = FileCell()
    cell.set_context(tmp_path, "Book")
    assert cell.attach(str(f))
    cell.materialize()
    files = json.loads(cell.source())["files"]
    assert len(files) == 1
    assert files[0]["name"] == "data.csv"
    assert "embed" in files[0] and "path" not in files[0]


def test_file_cell_large_file_goes_to_sidecar(qapp, tmp_path):
    from khervebook.filecell import FileCell, EMBED_LIMIT
    f = tmp_path / "big.bin"
    f.write_bytes(b"X" * (EMBED_LIMIT + 1024))
    cell = FileCell()
    cell.set_context(tmp_path, "Book")
    cell.attach(str(f))
    cell.materialize()
    files = json.loads(cell.source())["files"]
    assert files[0].get("path") == "Book_files/big.bin"
    assert "embed" not in files[0]
    assert (tmp_path / "Book_files" / "big.bin").exists()


def test_file_cell_holds_multiple_files(qapp, tmp_path):
    from khervebook.notebook import NotebookWidget
    (tmp_path / "a.txt").write_text("alpha")
    (tmp_path / "b.txt").write_text("beta")
    book = tmp_path / "Book.kbook"
    nb = NotebookWidget()
    nb.set_document_path(str(book))
    fc = nb.add_cell("file")
    fc.attach(str(tmp_path / "a.txt"))
    fc.attach(str(tmp_path / "b.txt"))
    assert fc.file_names == ["a.txt", "b.txt"]
    nb.prepare_save()
    book.write_text(nb.to_json())

    reopened = NotebookWidget()
    reopened.load_json(book.read_text())
    reopened.set_document_path(str(book))
    assert reopened._resolve_file("a.txt")
    assert reopened._resolve_file("b.txt")
    r = reopened.kernel.run("open(kf('b.txt')).read()")
    assert not r.error and r.result_repr == "'beta'"


def test_file_cell_reads_legacy_single_file_format(qapp):
    import base64
    from khervebook.filecell import FileCell
    legacy = json.dumps({"kbook_file": 1, "name": "old.txt", "size": 3,
                         "embed": base64.b64encode(b"abc").decode()})
    cell = FileCell()
    cell.set_source(legacy)
    assert cell.file_names == ["old.txt"]


def test_file_cell_text_preview_non_none(qapp, tmp_path):
    from khervebook.filecell import FileCell
    cell = FileCell()
    assert cell._as_text(b"a,b\n1,2\n", ".csv") is not None
    assert cell._as_text(b"\x00\x01\x02\xff", ".bin") is None


# -- structured-format previews -------------------------------------------

def test_preview_xlsx():
    import io
    from openpyxl import Workbook
    from khervebook import filepreview
    wb = Workbook()
    wb.active.title = "Data"
    wb.active.append(["name", "value"])
    wb.active.append(["Alpha", 3.14])
    wb.create_sheet("Notes")
    b = io.BytesIO()
    wb.save(b)
    summary = filepreview.summarize("x.xlsx", b.getvalue())
    assert summary and "Excel workbook" in summary
    assert "Data" in summary and "Notes" in summary
    assert "Alpha" in summary


def test_preview_ksheet():
    import io
    import numpy as np
    import h5py
    from khervebook import filepreview
    bio = io.BytesIO()
    with h5py.File(bio, "w") as f:
        f.attrs["format"] = "ksheet"
        f.attrs["sheet_count"] = 2
        g = f.create_group("sheet_0")
        g.attrs["name"] = "Survey"
        g.attrs["rows"] = 100
        g.attrs["cols"] = 5
        g.create_dataset("cells", data=np.array([b""] * 500))
        g2 = f.create_group("sheet_1")
        g2.attrs["name"] = "Fit"
        g2.attrs["rows"] = 20
        g2.attrs["cols"] = 3
        g2.create_dataset("cells", data=np.array([b""] * 60))
    summary = filepreview.summarize("wb.ksheet", bio.getvalue())
    assert summary and "KherveSheet workbook" in summary
    assert "Survey (100×5)" in summary and "Fit (20×3)" in summary


def test_preview_kfit():
    import io
    import numpy as np
    import h5py
    from khervebook import filepreview
    bio = io.BytesIO()
    with h5py.File(bio, "w") as f:
        f.attrs["format"] = "kfitting"
        cl = f.create_group("core_levels")
        for i, nm in enumerate(["Survey", "C1s", "O1s"]):
            g = cl.create_group(f"cl_{i}")
            g.attrs["name"] = nm
    summary = filepreview.summarize("p.kfit", bio.getvalue())
    assert summary and "KherveFitting project" in summary
    assert "3 core levels" in summary
    assert "Survey" in summary and "C1s" in summary


def test_file_cell_shows_structured_preview(qapp, tmp_path):
    import io
    from openpyxl import Workbook
    from khervebook.filecell import FileCell
    wb = Workbook()
    wb.active.append(["a", "b"])
    p = tmp_path / "book.xlsx"
    wb.save(str(p))
    cell = FileCell()
    cell.attach(str(p))
    # a preview widget (not the generic binary note) is produced
    att = cell._items[0]
    w = cell._preview_widget(att)
    from PyQt5.QtWidgets import QLabel
    assert isinstance(w, QLabel)
    assert "Excel workbook" in w.text()


def test_file_cell_resolved_path_and_kf(qapp, tmp_path):
    from khervebook.notebook import NotebookWidget
    small = tmp_path / "notes.txt"
    small.write_text("hello attach")
    book = tmp_path / "Book.kbook"
    nb = NotebookWidget()
    nb.set_document_path(str(book))
    fc = nb.add_cell("file")
    fc.attach(str(small))
    nb.prepare_save()
    book.write_text(nb.to_json())

    reopened = NotebookWidget()
    reopened.load_json(book.read_text())
    reopened.set_document_path(str(book))
    res = reopened.kernel.run("open(kf('notes.txt')).read()")
    assert not res.error
    assert res.result_repr == "'hello attach'"


def test_kf_no_arg_returns_notebook_dir(qapp, tmp_path):
    from khervebook.notebook import NotebookWidget
    book = tmp_path / "Book.kbook"
    nb = NotebookWidget()
    nb.set_document_path(str(book))
    res = nb.kernel.run("import os; os.path.realpath(kf())")
    assert not res.error
    assert os_samefile(res.result_repr, tmp_path)


def os_samefile(result_repr, tmp_path):
    import ast
    from pathlib import Path
    return Path(ast.literal_eval(result_repr)).resolve() == tmp_path.resolve()


def test_file_cell_drop_attaches(qapp, tmp_path):
    from khervebook.notebook import NotebookWidget
    f = tmp_path / "thing.xyz"          # unrecognised extension
    f.write_bytes(b"payload")
    nb = NotebookWidget()
    fc = nb.add_cell("file")
    assert nb.open_file_in_cell(str(f), fc)
    assert fc.file_name == "thing.xyz"
